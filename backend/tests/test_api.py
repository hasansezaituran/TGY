"""Gerçek .xlsx ve izole SQLite ile bağımsız API kabul testleri."""
import importlib
import io
import time
from contextlib import contextmanager
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine

# Kaynak başlıkları uygulama kodundan kopyalanarak değil kullanıcı sözleşmesinden tanımlanır.
SHEETS = {
    'orders': ('Sipariş Kalemleri', ['Malzeme Kodu', 'Malzeme İsmi', 'Proje Kodu', 'Tedarikçi Numarası', 'Tedarikçi İsmi', 'Sipariş Numarası', 'Sipariş Kalem', 'Miktar', 'Gelen Miktar', 'İade Miktar', 'Teslim Durumu', 'Birim Fiyatı', 'Teslimata Açık Kapalılık Durumu', 'İstek Açık Kapalı Durumu', 'Sipariş Durumu', 'Sipariş Onay Durumu', 'Para Birimi'], ['material', 'name', 'project', 'supplier', 'supplierName', 'number', 'line', 'qty', 'received', 'returned', 'delivery', 'price', 'deliveryOpen', 'requestOpen', 'status', 'approval', 'currency']),
    'parts': ('Parça Bilgileri', ['Malzeme Kodu', 'Eski Malzeme Kodu', 'Malzeme İsmi', 'Üretim Durumu (Aktif/Pasif)', 'MRP Sorumlusu', 'Üretim Adam Saati', 'Trim Adam Saati'], ['id', 'old', 'name', 'status', 'mrp', 'hours', 'trim']),
    'movements': ('Kalıp Bilgileri', ['Malzeme Kodu', 'Malzeme İsmi', 'Kalıp ATA Kodu', 'Kalıp ASY Kodu', 'Kaynak Depo', 'Hedef Depo', 'Sevk Tarihi', 'Sevk Sebebi'], ['material', 'name', 'ata', 'asy', 'from', 'to', 'date', 'reason']),
    'shipments': ('Sevkiyat Listesi', ['Tedarikçi', 'Sevk Tarihi', 'Malzeme Kodu', 'Malzeme İsmi', 'Seri Numarası', 'FAİ Seri Parça Durumu', 'Trimli Trimsiz Durumu', 'Sipariş Numarası', 'İrsaliye Numarası'], ['supplier', 'date', 'material', 'name', 'serial', 'fai', 'trim', 'order', 'note']),
}


def sample():
    return {
        'orders': [{'id': 'o1', 'material': '0001', 'name': 'Braket', 'project': 'P2027', 'supplier': '00100', 'supplierName': 'A Firma', 'number': '450001', 'line': '010', 'qty': 120, 'received': 120, 'returned': 5, 'delivery': 'Kısmi', 'price': 12.5, 'deliveryOpen': 'Açık', 'requestOpen': 'Açık', 'status': 'Açık', 'approval': 'Onaylandı', 'currency': 'USD'}],
        'parts': [{'id': '0001', 'old': 'ESK-001', 'name': 'Braket', 'status': 'Aktif', 'mrp': 'Deniz', 'hours': 1.5, 'trim': 0.25}],
        'movements': [{'id': 'm1', 'material': '0001', 'name': 'Braket', 'ata': 'K001', 'asy': 'ASY001', 'from': 'MERKEZ', 'to': 'A DEPO', 'date': '2026-12-31', 'reason': 'Üretim'}],
        'shipments': [{'id': 's1', 'supplier': '00100', 'date': '2026-12-31', 'material': '0001', 'name': 'Braket', 'serial': 'SN001', 'fai': 'Seri', 'trim': 'Trimli', 'order': '450001', 'note': 'IRS001'}],
    }


@pytest.fixture
def api_db(tmp_path, monkeypatch):
    storage = importlib.import_module('backend.storage')
    api = importlib.import_module('backend.main')
    engines = []

    @contextmanager
    def open_client(filename='test.db'):
        engine = create_engine(f'sqlite:///{tmp_path / filename}', connect_args={'check_same_thread': False})
        engines.append(engine)
        monkeypatch.setattr(storage, 'engine', engine)
        if hasattr(api, 'engine'):
            monkeypatch.setattr(api, 'engine', engine)
        storage.initialize()
        with TestClient(api.app) as client:
            yield client
        engine.dispose()

    yield open_client
    for engine in engines:
        engine.dispose()


def state(client):
    response = client.get('/api/state')
    assert response.status_code == 200, response.text
    return response.json()


def write(client, changes, revision=None):
    if revision is None:
        revision = state(client)['revision']
    return client.post('/api/changes', json={'expected_revision': revision, 'changes': changes})


def source_changes(data):
    return [{'entity': entity, 'value': row} for entity, rows in data.items() for row in rows]


def workbook(data):
    book = Workbook(write_only=True)
    for entity, rows in data.items():
        name, headers, keys = SHEETS[entity]
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key, '') for key in keys])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def preview(client, content, kind='auto', filename='veri.xlsx'):
    response = client.post('/api/imports/preview', files={'file': (filename, content, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, data={'kind': kind})
    assert response.status_code == 200, response.text
    return response.json()


def commit(client, report, **kwargs):
    return client.post(f"/api/imports/{report['id']}/commit", json={'expected_revision': report['revision'], **kwargs})


def test_empty_real_database_has_confirmed_settings(api_db):
    with api_db() as client:
        current = state(client)
        assert current['revision'] == 0
        assert all(current['data'][entity] == [] for entity in SHEETS)
        assert current['data']['settings'] == {'returns': 'gross', 'molds': 'any', 'shipmentUnits': True, 'hoursUnit': 'piece'}


def test_atomic_upsert_preserves_other_entities_and_survives_restart(api_db):
    with api_db() as client:
        changes = source_changes(sample()) + [
            {'entity': 'requirements', 'id': '0001', 'value': ['K001']},
            {'entity': 'mappings', 'id': '00100', 'value': ['A DEPO']},
            {'entity': 'links', 'id': 's1', 'value': 'o1'},
        ]
        response = write(client, changes)
        assert response.status_code == 200, response.text
        first = state(client)
        assert first['revision'] == 1
        assert first['data']['orders'][0]['received'] == 120
        assert first['data']['links']['s1'] == 'o1'
    with api_db() as client:
        assert state(client) == first
        part = {**sample()['parts'][0], 'status': 'Pasif'}
        assert write(client, [{'entity': 'parts', 'value': part}]).status_code == 200
        current = state(client)
        assert current['data']['orders'] == first['data']['orders']
        assert current['data']['shipments'] == first['data']['shipments']
        assert current['data']['movements'] == first['data']['movements']
        assert current['data']['parts'][0]['status'] == 'Pasif'


def test_stale_revision_returns_409_without_overwrite(api_db):
    with api_db() as client:
        assert write(client, [{'entity': 'parts', 'value': sample()['parts'][0]}], revision=0).status_code == 200
        response = write(client, [{'entity': 'parts', 'value': {**sample()['parts'][0], 'status': 'Pasif'}}], revision=0)
        assert response.status_code == 409, response.text
        assert state(client)['data']['parts'][0]['status'] == 'Aktif'
        assert state(client)['revision'] == 1


def test_invalid_change_rolls_back_whole_batch(api_db):
    with api_db() as client:
        before = state(client)
        data = sample()
        data['orders'][0]['qty'] = -1
        response = write(client, source_changes(data))
        assert response.status_code in (400, 422), response.text
        assert state(client) == before


def test_natural_order_identity_conflict_is_atomic(api_db):
    with api_db() as client:
        order = sample()['orders'][0]
        response = write(client, [{'entity': 'orders', 'value': order}, {'entity': 'orders', 'value': {**order, 'id': 'o2', 'material': 'OTHER'}}])
        assert response.status_code in (400, 409, 422), response.text
        assert state(client)['data']['orders'] == []
        assert state(client)['revision'] == 0


def test_excel_preview_does_not_change_sources_then_commit_preserves_past_and_codes(api_db):
    with api_db() as client:
        before = state(client)
        report = preview(client, workbook(sample()))
        assert report['can_commit'] is True, report
        assert report['errors'] == []
        assert sum(group['new'] for group in report['groups']) == 4
        assert state(client) == before
        response = commit(client, report)
        assert response.status_code == 200, response.text
        current = state(client)
        assert current['revision'] == 1
        assert all(len(current['data'][entity]) == 1 for entity in SHEETS)
        assert current['data']['orders'][0]['material'] == '0001'
        assert current['data']['orders'][0]['line'] == '010'
        assert current['data']['shipments'][0]['date'] == '2026-12-31'
        assert current['data']['orders'][0]['currency'] == 'USD'


def test_default_import_preserves_existing_explicit_update_changes_target_only(api_db):
    with api_db() as client:
        assert write(client, source_changes(sample())).status_code == 200
        altered = sample()['parts'][0] | {'status': 'Pasif'}
        report = preview(client, workbook({'parts': [altered]}))
        assert report['groups'][0]['existing'] == 1
        assert commit(client, report).status_code == 200
        assert state(client)['data']['parts'][0]['status'] == 'Aktif'
        report = preview(client, workbook({'parts': [altered]}))
        assert commit(client, report, update_existing=True).status_code == 200
        current = state(client)
        assert current['data']['parts'][0]['status'] == 'Pasif'
        assert current['data']['orders'] == sample()['orders']
        assert current['data']['shipments'] == sample()['shipments']


def test_stale_preview_commit_does_not_overwrite_new_api_changes(api_db):
    with api_db() as client:
        report = preview(client, workbook(sample()))
        assert write(client, [{'entity': 'parts', 'value': sample()['parts'][0] | {'status': 'Pasif'}}]).status_code == 200
        before = state(client)
        response = commit(client, report)
        assert response.status_code == 409, response.text
        assert state(client) == before


def test_identical_duplicate_rows_reported_and_committed_once(api_db):
    with api_db() as client:
        row = sample()['parts'][0]
        report = preview(client, workbook({'parts': [row, row]}))
        assert report['can_commit'] is True
        assert report['groups'][0]['duplicate'] == 1
        assert report['groups'][0]['new'] == 1
        assert commit(client, report).status_code == 200
        assert len(state(client)['data']['parts']) == 1


def test_conflicting_duplicate_rows_cannot_commit(api_db):
    with api_db() as client:
        row = sample()['parts'][0]
        before = state(client)
        report = preview(client, workbook({'parts': [row, row | {'status': 'Pasif'}]}))
        assert report['can_commit'] is False
        assert report['errors']
        assert commit(client, report).status_code in (400, 422)
        assert state(client) == before


@pytest.mark.parametrize('bad_date', ['2027-02-30', '2027-13-01', 'tarih-yok'])
def test_invalid_excel_dates_are_reported_without_partial_sources(api_db, bad_date):
    with api_db() as client:
        data = sample()
        data['shipments'][0]['date'] = bad_date
        before = state(client)
        report = preview(client, workbook(data))
        assert report['can_commit'] is False
        assert any(error['field'] == 'Sevk Tarihi' and error['row'] == 2 for error in report['errors'])
        assert state(client) == before


def test_excel_native_datetime_is_accepted(api_db):
    with api_db() as client:
        data = sample()
        data['movements'][0]['date'] = datetime(2026, 12, 31)
        report = preview(client, workbook(data))
        assert report['can_commit'] is True, report
        assert commit(client, report).status_code == 200
        assert state(client)['data']['movements'][0]['date'] == '2026-12-31'


def test_missing_header_and_non_numeric_values_have_field_errors(api_db):
    with api_db() as client:
        data = sample()
        data['orders'][0]['qty'] = 'bir miktar'
        report = preview(client, workbook({'orders': data['orders']}))
        assert report['can_commit'] is False
        assert any(error['field'] == 'Miktar' for error in report['errors'])
        book = Workbook()
        book.active.title = 'Parça Bilgileri'
        book.active.append(['Malzeme Kodu'])
        book.active.append(['0001'])
        stream = io.BytesIO()
        book.save(stream)
        report = preview(client, stream.getvalue())
        assert report['can_commit'] is False
        assert any(error['row'] == 1 and 'Malzeme İsmi' in error['field'] for error in report['errors'])
        assert state(client)['revision'] == 0


def test_malformed_workbook_rejected_without_server_error(api_db):
    with api_db() as client:
        before = state(client)
        response = client.post('/api/imports/preview', files={'file': ('bozuk.xlsx', b'not-an-xlsx', 'application/octet-stream')}, data={'kind': 'auto'})
        assert response.status_code in (400, 422), response.text
        assert state(client) == before


def test_export_template_headers_and_source_roundtrip(api_db):
    with api_db('source.db') as client:
        assert write(client, source_changes(sample())).status_code == 200
        template = client.get('/api/template')
        assert template.status_code == 200
        book = load_workbook(io.BytesIO(template.content), read_only=True)
        for _, (title, headers, _) in SHEETS.items():
            assert list(next(book[title].values)) == headers
        book.close()
        response = client.get('/api/export')
        assert response.status_code == 200
        exported = response.content
    with api_db('target.db') as client:
        report = preview(client, exported)
        assert report['can_commit'] is True, report
        assert commit(client, report).status_code == 200
        restored = state(client)['data']
        for entity, rows in sample().items():
            actual = restored[entity][0]
            for field, value in rows[0].items():
                if field != 'id' or entity == 'parts':
                    assert actual[field] == value, (entity, field, actual[field], value)


def test_40000_real_xlsx_rows_preview_commit_and_reopen(api_db):
    rows = []
    row = sample()['shipments'][0]
    for index in range(40_000):
        rows.append(row | {'id': f's{index}', 'serial': f'SN-{index:05}', 'note': f'IRS-{index:05}'})
    content = workbook({'shipments': rows})
    started = time.perf_counter()
    with api_db() as client:
        report = preview(client, content)
        preview_seconds = time.perf_counter() - started
        assert report['can_commit'] is True, report
        assert report['groups'][0]['rows'] == 40_000
        assert report['groups'][0]['new'] == 40_000
        assert len(report['groups'][0]['preview']) <= 20
        assert state(client)['data']['shipments'] == []
        response = commit(client, report)
        assert response.status_code == 200, response.text
        assert len(state(client)['data']['shipments']) == 40_000
    with api_db() as client:
        current = state(client)
        assert len(current['data']['shipments']) == 40_000
        assert any(row['serial'] == 'SN-39999' for row in current['data']['shipments'])
    print({'senaryo': 'gercek-xlsx-sentetik-veri', 'satir': 40_000, 'dosya_bayt': len(content), 'onizleme_saniye': round(preview_seconds, 3), 'toplam_saniye': round(time.perf_counter() - started, 3)})


def test_invalid_manual_link_rejected_atomically(api_db):
    with api_db() as client:
        assert write(client, source_changes(sample())).status_code == 200
        before = state(client)
        response = write(client, [{'entity': 'links', 'id': 's1', 'value': 'olmayan-siparis'}])
        assert response.status_code == 422
        assert state(client) == before


def test_import_update_cannot_invalidate_existing_confirmed_link(api_db):
    from sqlalchemy.orm import Session
    from backend import storage
    with api_db() as client:
        changes = source_changes(sample()) + [{'entity': 'links', 'id': 's1', 'value': 'o1'}]
        assert write(client, changes).status_code == 200
        before = state(client)
        before_audit = client.get('/api/audit').json()['items']
        altered = sample()['orders'][0] | {'material': 'FARKLI-PARCA'}
        report = preview(client, workbook({'orders': [altered]}))
        assert report['can_commit'] is True
        response = commit(client, report, update_existing=True)
        assert response.status_code == 422, response.text
        assert state(client) == before
        assert client.get('/api/audit').json()['items'] == before_audit
        with Session(storage.engine) as session:
            assert session.get(storage.ImportBatch, report['id']).status == 'preview'


def test_corrupt_xml_in_valid_zip_rejected_without_server_error(api_db):
    import zipfile
    valid = workbook({'parts': sample()['parts']})
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(valid)) as source, zipfile.ZipFile(target, 'w') as destination:
        for member in source.infolist():
            destination.writestr(member, b'<gecersiz-xml' if member.filename == 'xl/workbook.xml' else source.read(member.filename))
    with api_db() as client:
        before = state(client)
        response = client.post('/api/imports/preview', files={'file': ('bozuk.xlsx', target.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, data={'kind': 'auto'})
        assert response.status_code in (400, 422), response.text
        assert state(client) == before


def test_computed_fields_cannot_be_written_as_source(api_db):
    with api_db() as client:
        before = state(client)
        row = sample()['orders'][0] | {'state': 'Üretilebilir', 'net': 999}
        response = write(client, [{'entity': 'orders', 'value': row}])
        assert response.status_code == 422
        assert state(client) == before


def test_python_assessments_and_paged_orders_use_confirmed_rules(api_db):
    with api_db() as client:
        changes = source_changes(sample()) + [
            {'entity': 'requirements', 'id': '0001', 'value': ['K001', 'K002']},
            {'entity': 'mappings', 'id': '00100', 'value': ['A DEPO']},
            {'entity': 'links', 'id': 's1', 'value': 'o1'},
        ]
        assert write(client, changes).status_code == 200
        response = client.get('/api/assessments')
        assert response.status_code == 200
        body = response.json()
        assert body['revision'] == 1
        value = body['items']['o1']
        assert value['state'] == 'Üretilebilir'
        assert value['net'] == 115 and value['remaining'] == 5
        assert value['shipments'] == ['s1']
        response = client.get('/api/orders', params={'q': '0001', 'page_size': 1})
        assert response.status_code == 200
        assert response.json()['total'] == 1
        assert response.json()['items'][0]['computed']['net'] == 115


def test_import_commit_cannot_be_applied_twice(api_db):
    with api_db() as client:
        report = preview(client, workbook({'parts': sample()['parts']}))
        assert commit(client, report).status_code == 200
        before = state(client)
        assert commit(client, report).status_code == 409
        assert state(client) == before


def test_unsupported_delete_action_does_not_remove_records(api_db):
    with api_db() as client:
        assert write(client, source_changes(sample())).status_code == 200
        before = state(client)
        response = write(client, [{'entity': 'parts', 'id': '0001', 'value': sample()['parts'][0], 'delete': True}])
        assert response.status_code == 422
        assert state(client) == before


def test_original_import_source_bytes_are_preserved_and_unknown_source_is_404(api_db):
    with api_db() as client:
        content = workbook(sample())
        report = preview(client, content, filename='orijinal kaynak.xlsx')
        assert report['can_commit'] is True
        response = client.get(f"/api/imports/{report['id']}/source")
        assert response.status_code == 200
        assert response.content == content
        assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        assert commit(client, report).status_code == 200
    with api_db() as client:
        response = client.get(f"/api/imports/{report['id']}/source")
        assert response.status_code == 200
        assert response.content == content
        assert client.get('/api/imports/olmayan-aktarim/source').status_code == 404
