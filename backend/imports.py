"""Excel doğrulama, güvenli önizleme ve kaynak çalışma kitabı üretimi."""
import hashlib
import io
import json
import math
import re
import zipfile
from xml.etree.ElementTree import ParseError
from collections import Counter, defaultdict
from datetime import date, datetime
from uuid import uuid4
import pandas as pd
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from .models import MODELS
from .storage import natural_key

SHEETS = {
    "orders": ("Sipariş Kalemleri", [
        ("Malzeme Kodu", "material"), ("Malzeme İsmi", "name"), ("Proje Kodu", "project"),
        ("Tedarikçi Numarası", "supplier"), ("Tedarikçi İsmi", "supplierName"), ("Sipariş Numarası", "number"),
        ("Sipariş Kalem", "line"), ("Miktar", "qty"), ("Gelen Miktar", "received"), ("İade Miktar", "returned"),
        ("Teslim Durumu", "delivery"), ("Birim Fiyatı", "price"), ("Teslimata Açık Kapalılık Durumu", "deliveryOpen"),
        ("İstek Açık Kapalı Durumu", "requestOpen"), ("Sipariş Durumu", "status"), ("Sipariş Onay Durumu", "approval")]),
    "movements": ("Kalıp Bilgileri", [
        ("Malzeme Kodu", "material"), ("Malzeme İsmi", "name"), ("Kalıp ATA Kodu", "ata"), ("Kalıp ASY Kodu", "asy"),
        ("Kaynak Depo", "from"), ("Hedef Depo", "to"), ("Sevk Tarihi", "date"), ("Sevk Sebebi", "reason")]),
    "parts": ("Parça Bilgileri", [
        ("Malzeme Kodu", "id"), ("Eski Malzeme Kodu", "old"), ("Malzeme İsmi", "name"), ("Üretim Durumu (Aktif/Pasif)", "status"),
        ("MRP Sorumlusu", "mrp"), ("Üretim Adam Saati", "hours"), ("Trim Adam Saati", "trim")]),
    "shipments": ("Sevkiyat Listesi", [
        ("Tedarikçi", "supplier"), ("Sevk Tarihi", "date"), ("Malzeme Kodu", "material"), ("Malzeme İsmi", "name"),
        ("Seri Numarası", "serial"), ("FAİ Seri Parça Durumu", "fai"), ("Trimli Trimsiz Durumu", "trim"),
        ("Sipariş Numarası", "order"), ("İrsaliye Numarası", "note")])}
NUMERIC = {"qty", "received", "returned", "price", "hours", "trim"}
IDENTIFIERS = {"id", "old", "material", "supplier", "number", "line", "ata", "asy", "serial", "order", "note"}
MAX_ROWS = 100000
MAX_BYTES = 16 * 1024 * 1024
REPORT_LIMIT = 200


class WorkbookError(ValueError):
    pass


def make_workbook(data=None):
    workbook = Workbook(write_only=True)
    for entity, (name, fields) in SHEETS.items():
        sheet = workbook.create_sheet(name)
        columns = fields + ([("Para Birimi", "currency")] if entity == "orders" else [])
        sheet.append([label for label, _ in columns])
        for row in (data or {}).get(entity, []):
            # Formül enjeksiyonunu önlemek için tüm metinleri açıkça metin hücresi yaz.
            from openpyxl.cell import WriteOnlyCell
            cells = []
            for _, key in columns:
                value = row.get(key, "")
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def preview(content, kind, state):
    if len(content) > MAX_BYTES:
        raise WorkbookError("Dosya en fazla 16 MB olabilir")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
                raise WorkbookError("Açılmış çalışma kitabı boyutu sınırı aşıyor")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, ParseError) as error:
        raise WorkbookError("Geçerli bir .xlsx çalışma kitabı okunamadı") from error
    errors, warnings, groups = [], [], []
    totals = {"errors": 0, "warnings": 0}

    def report(target, sheet, row, field, message):
        key = "errors" if target is errors else "warnings"
        totals[key] += 1
        if len(target) < REPORT_LIMIT:
            target.append({"sheet": sheet, "row": row, "field": field, "message": message})

    known_names = defaultdict(set)
    known_ids = set()
    for order in state["data"]["orders"]:
        known_names[order["supplierName"]].add(order["supplier"])
        known_ids.add(order["supplier"])
    staged = defaultdict(list)
    existing = {entity: {natural_key(entity, row): row for row in state["data"][entity]} for entity in MODELS}
    seen = {entity: {} for entity in MODELS}
    total_rows = 0
    recognized = 0
    selected = list(workbook.worksheets)
    if kind != "auto" and len(selected) != 1:
        workbook.close()
        raise WorkbookError("Belirli liste türü seçildiğinde dosya tek çalışma sayfası içermelidir")
    # Siparişleri önce oku: aynı dosyadaki sevkiyatların firma adları çözülebilsin.
    selected.sort(key=lambda sheet: 0 if sheet.title == SHEETS["orders"][0] or kind == "orders" else 1)
    for sheet in selected:
        if (sheet.max_row and sheet.max_row > MAX_ROWS + 1) or (sheet.max_column and sheet.max_column > 256):
            workbook.close()
            raise WorkbookError("Çalışma sayfası boyutları güvenli satır/sütun sınırını aşıyor")
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        headers = [cell.value for cell in header_cells]
        entity = kind if kind != "auto" else next((key for key, (name, _) in SHEETS.items() if sheet.title == name), None)
        if entity is None:
            report(warnings, sheet.title, 1, "", "Çalışma sayfası adı tanınmadı; bu sayfa aktarılmayacak")
            continue
        recognized += 1
        name, fields = SHEETS[entity]
        if len(headers) != len(set(str(h) for h in headers if h is not None)):
            # Sondaki biçimlendirilmiş boş hücreleri başlık saymayız.
            present = [h for h in headers if h is not None]
            if len(present) != len(set(map(str, present))):
                report(errors, sheet.title, 1, "", "Mükerrer sütun başlığı")
                continue
        positions = {label: index for index, label in enumerate(headers) if isinstance(label, str)}
        missing = [label for label, _ in fields if label not in positions]
        for label in missing:
            report(errors, sheet.title, 1, label, "Zorunlu kaynak sütunu eksik")
        if missing:
            continue
        if entity == "orders" and "Para Birimi" not in positions:
            report(warnings, sheet.title, 1, "Para Birimi", "Kullanıcının doğruladığı USD para birimi uygulanacak")
        rows_count = duplicate = new = updated = 0
        for row_number, cells in enumerate(iterator, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            total_rows += 1
            rows_count += 1
            if total_rows > MAX_ROWS:
                workbook.close()
                raise WorkbookError("Çalışma kitabı toplam 100.000 veri satırını aşamaz")
            row = {}
            invalid = False
            columns = fields + ([("Para Birimi", "currency")] if "Para Birimi" in positions and entity == "orders" else [])
            for label, key in columns:
                cell = cells[positions[label]] if positions[label] < len(cells) else None
                value = cell.value if cell else None
                if cell and cell.data_type == "f":
                    report(errors, sheet.title, row_number, label, "Formül hücresi kabul edilmez; doğrulanmış değeri yapıştırın")
                    invalid = True
                    continue
                if key in NUMERIC and not (entity == "shipments" and key == "trim"):
                    if isinstance(value, bool) or value is None or isinstance(value, str) and not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", value):
                        report(errors, sheet.title, row_number, label, "Geçerli sayı gerekli; ondalık ayracı metinde nokta olmalıdır")
                        invalid = True
                        continue
                    try:
                        number = float(value)
                        if not math.isfinite(number) or number < 0:
                            raise ValueError()
                        row[key] = number
                    except (ValueError, TypeError):
                        report(errors, sheet.title, row_number, label, "Negatif veya geçersiz sayı")
                        invalid = True
                elif key == "date":
                    try:
                        if isinstance(value, datetime):
                            if value.time().isoformat() != "00:00:00":
                                raise ValueError()
                            parsed = value.date()
                        elif isinstance(value, date):
                            parsed = value
                        elif isinstance(value, str):
                            parsed = date.fromisoformat(value) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) else datetime.strptime(value, "%d.%m.%Y").date()
                        else:
                            raise ValueError()
                        row[key] = parsed.isoformat()
                    except (ValueError, TypeError):
                        report(errors, sheet.title, row_number, label, "Geçerli tarih gerekli: GG.AA.YYYY veya YYYY-AA-GG")
                        invalid = True
                elif value is None:
                    row[key] = ""
                elif isinstance(value, str):
                    row[key] = value
                elif key in IDENTIFIERS and isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value == int(value):
                    row[key] = str(int(value))
                    report(warnings, sheet.title, row_number, label, "Sayısal kimlik metne çevrildi; Excel'de kaybolmuş baştaki sıfırlar geri getirilemez")
                else:
                    report(errors, sheet.title, row_number, label, "Metin hücresi gerekli")
                    invalid = True
            if invalid:
                continue
            if entity == "orders":
                row.setdefault("currency", "USD")
            if entity == "shipments":
                supplier = row["supplier"]
                if supplier in known_ids:
                    pass
                elif len(known_names[supplier]) == 1:
                    row["supplier"] = next(iter(known_names[supplier]))
                    report(warnings, sheet.title, row_number, "Tedarikçi", "Tedarikçi adı bilinen tek bir numarayla birebir eşleştirildi")
                else:
                    report(warnings, sheet.title, row_number, "Tedarikçi", "Tedarikçi çözümlenemedi; kaynak değer korunacak, kesin sipariş eşleşmesi yapılmayabilir")
                if not row["serial"]:
                    report(warnings, sheet.title, row_number, "Seri Numarası", "Seri numarası eksik")
                if not row["note"]:
                    report(warnings, sheet.title, row_number, "İrsaliye Numarası", "İrsaliye numarası eksik")
            if entity != "parts":
                row["id"] = ""
                key = natural_key(entity, row)
                row["id"] = hashlib.sha256((entity + key).encode()).hexdigest()
            key = natural_key(entity, row)
            if key in existing[entity]:
                row["id"] = existing[entity][key]["id"]
            try:
                row = MODELS[entity].model_validate(row).model_dump(by_alias=True)
            except ValidationError as error:
                for issue in error.errors():
                    report(errors, sheet.title, row_number, str(issue["loc"][0]), "Geçersiz veya eksik kaynak alanı: " + issue["msg"])
                continue
            if key in seen[entity]:
                duplicate += 1
                if seen[entity][key] != row:
                    report(errors, sheet.title, row_number, "", "Aynı kimlikle çelişkili birden fazla satır var")
                continue
            seen[entity][key] = row
            if key in existing[entity]:
                updated += 1
            else:
                new += 1
            staged[entity].append(row)
            if entity == "orders":
                known_ids.add(row["supplier"])
                known_names[row["supplierName"]].add(row["supplier"])
        groups.append({"entity": entity, "sheet": sheet.title, "rows": rows_count, "new": new, "existing": updated, "duplicate": duplicate, "preview": staged[entity][:20]})
    workbook.close()
    if not recognized:
        report(errors, "", 1, "", "Tanınan çalışma sayfası bulunamadı; şablon kullanın veya tek sayfa türünü seçin")
    if not total_rows and not totals["errors"]:
        report(errors, "", 1, "", "Aktarılacak veri satırı bulunamadı")
    # Pandas ile toplu seri tekrar analizi; satır başına tüm sevkiyatı taramaz.
    shipments = staged.get("shipments", [])
    if shipments:
        frame = pd.DataFrame(shipments)
        counts = Counter(row["serial"] for row in state["data"]["shipments"] if row["serial"])
        for serial, count in frame.loc[frame["serial"] != "", "serial"].value_counts().items():
            if count > 1 or counts[serial] > 0:
                report(warnings, SHEETS["shipments"][0], 0, "Seri Numarası", f"Mükerrer seri kontrolü gerekli: {serial}")
    report_data = {"id": str(uuid4()), "revision": state["revision"], "can_commit": totals["errors"] == 0, "errors": errors, "warnings": warnings, "error_count": totals["errors"], "warning_count": totals["warnings"], "reports_truncated": totals["errors"] > len(errors) or totals["warnings"] > len(warnings), "groups": groups}
    return report_data, dict(staged)
